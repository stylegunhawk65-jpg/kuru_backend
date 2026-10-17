"""Unit and integration tests for Demo 100% Match, Notification Actions, and LinkedIn verification."""

from unittest.mock import patch
import pytest
from httpx import AsyncClient

from app.services.demo_match_service import demo_match_service

DEMO_USER_UID = "usr_demo_tester"
DEMO_HEADERS = {"Authorization": "Bearer token_demo_tester"}


def mock_verify_id_token(token: str, *args, **kwargs):
    if token == "token_demo_tester":
        return {
            "uid": DEMO_USER_UID,
            "email": "tester@kuru.ai",
            "name": "Demo Tester",
        }
    raise ValueError("Invalid test token")


@pytest.fixture(autouse=True)
def patch_firebase_auth():
    with patch("app.core.auth.firebase_auth.verify_id_token", side_effect=mock_verify_id_token):
        yield


@pytest.mark.asyncio
async def test_genuine_100_match_calculation():
    """Verify that the rule-based matcher produces genuine 100.0% without fabrication."""
    result = demo_match_service.compute_genuine_100_match()
    assert result.match_score == 100.0
    assert len(result.matched_skills) == 4
    assert len(result.missing_skills) == 0
    assert any("50.0 pts" in r for r in result.match_reasons)
    assert any("25.0 pts" in r for r in result.match_reasons)
    assert any("15.0 pts" in r for r in result.match_reasons)
    assert any("10.0 pts" in r for r in result.match_reasons)


@pytest.mark.asyncio
async def test_trigger_demo_match_idempotent(async_client: AsyncClient):
    """Verify trigger endpoint creates notification and repeated calls do not duplicate it."""
    # First trigger
    res1 = await async_client.post("/api/v1/jobs/demo/match-100", headers=DEMO_HEADERS)
    assert res1.status_code == 200
    data1 = res1.json()
    assert data1["match_score"] == 100.0
    assert data1["job"]["company"] == "DemoTech (DEMO)"
    assert data1["notification"]["data"]["is_demo"] is True
    notif_id_1 = data1["notification"]["id"]

    # Second trigger (Idempotency check)
    res2 = await async_client.post("/api/v1/jobs/demo/match-100", headers=DEMO_HEADERS)
    assert res2.status_code == 200
    data2 = res2.json()
    assert data2["notification"]["id"] == notif_id_1, "Repeated tap must reuse existing notification ID"


@pytest.mark.asyncio
async def test_demo_action_approve_workflow(async_client: AsyncClient):
    """Verify APPROVE_CV requires explicit confirmation and records approval against version."""
    # Trigger demo match first
    res = await async_client.post("/api/v1/jobs/demo/match-100", headers=DEMO_HEADERS)
    notif_id = res.json()["notification"]["id"]

    # Try to approve without review confirmation
    res_unconfirmed = await async_client.post(
        "/api/v1/jobs/demo/action",
        headers=DEMO_HEADERS,
        json={
            "notification_id": notif_id,
            "action": "APPROVE_CV",
            "confirmed_reviewed": False,
        },
    )
    assert res_unconfirmed.status_code == 200
    data_unconf = res_unconfirmed.json()
    assert data_unconf["status"] == "review_required"
    assert data_unconf["requires_dialog"] is True

    # Approve with explicit confirmation
    res_approved = await async_client.post(
        "/api/v1/jobs/demo/action",
        headers=DEMO_HEADERS,
        json={
            "notification_id": notif_id,
            "action": "APPROVE_CV",
            "confirmed_reviewed": True,
        },
    )
    assert res_approved.status_code == 200
    data_app = res_approved.json()
    assert data_app["status"] == "approved"
    assert "NOT submitted automatically" in data_app["message"]


@pytest.mark.asyncio
async def test_demo_action_reject(async_client: AsyncClient):
    """Verify REJECT marks notification dismissed without deleting real records."""
    res = await async_client.post("/api/v1/jobs/demo/match-100", headers=DEMO_HEADERS)
    notif_id = res.json()["notification"]["id"]

    res_reject = await async_client.post(
        "/api/v1/jobs/demo/action",
        headers=DEMO_HEADERS,
        json={
            "notification_id": notif_id,
            "action": "REJECT",
        },
    )
    assert res_reject.status_code == 200
    assert res_reject.json()["status"] == "rejected"


@pytest.mark.asyncio
async def test_demo_cv_update_flow(async_client: AsyncClient):
    """Verify demo CV can be viewed, updated, and version incremented to v1.1-demo."""
    # Get demo CV
    res_get = await async_client.get("/api/v1/jobs/demo/cv", headers=DEMO_HEADERS)
    assert res_get.status_code == 200
    assert res_get.json()["id"] == "demo-cv-100"

    # Update demo CV
    res_put = await async_client.put(
        "/api/v1/jobs/demo/cv",
        headers=DEMO_HEADERS,
        json={
            "summary": "Updated senior Python engineer profile with high concurrency expertise.",
            "skills": ["Python", "FastAPI", "PostgreSQL", "Docker", "AsyncIO"],
        },
    )
    assert res_put.status_code == 200
    updated_cv = res_put.json()
    assert updated_cv["version"] == "v1.1-demo"
    assert "AsyncIO" in updated_cv["skills"]


@pytest.mark.asyncio
async def test_linkedin_status_unconfigured(async_client: AsyncClient):
    """Verify LinkedIn OAuth url reports clean unconfigured status when server lacks credentials."""
    res = await async_client.get("/api/v1/profiles/linkedin/oauth/url", headers=DEMO_HEADERS)
    assert res.status_code == 200
    data = res.json()
    assert data["configured"] is False
    assert "LINKEDIN_CLIENT_ID" in data["missing_keys"]


@pytest.mark.asyncio
async def test_cv_upload_pdf_and_docx(async_client: AsyncClient):
    """Verify Master CV upload, parsing, and persistence for both PDF and DOCX documents."""
    import io
    import docx

    # 1. Test PDF Upload
    pdf_bytes = (
        b"%PDF-1.4\n"
        b"Candidate: Mayank Pawar\n"
        b"Email: mayank.pawar@example.com\n"
        b"Phone: +91 9876543210\n"
        b"EDUCATION: B.Tech Computer Science, Engineering Institute, 2026\n"
        b"EXPERIENCE: Software Developer Intern at Tech Labs (2024 - 2025)\n"
        b"Engineered scalable REST APIs in Python FastAPI and Flutter clients.\n"
        b"SKILLS: Python, FastAPI, Flutter, PostgreSQL, Docker, Git\n"
        b"%%EOF"
    )
    pdf_files = {"file": ("Mayank_CV.pdf", pdf_bytes, "application/pdf")}
    res_pdf = await async_client.post(
        "/api/v1/resumes/cv/upload",
        files=pdf_files,
        headers=DEMO_HEADERS,
    )
    assert res_pdf.status_code == 201
    pdf_data = res_pdf.json()
    assert pdf_data["file_type"] == "pdf"
    assert "Mayank" in (pdf_data.get("full_name") or "")
    assert "Python" in pdf_data.get("skills", [])

    # 2. Test DOCX Upload
    doc = docx.Document()
    doc.add_heading("Candidate: Mayank Pawar", 0)
    doc.add_paragraph("Email: mayank.docx@example.com | Phone: +91 9876543210")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("B.Tech Computer Science, Engineering Institute, 2026")
    doc.add_heading("EXPERIENCE", level=1)
    doc.add_paragraph("Software Engineer at Cloud Innovations (2024 - 2025)")
    doc.add_paragraph("Developed backend services using Python FastAPI and Docker.")
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL, Flutter")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    docx_files = {"file": ("Mayank_CV.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res_docx = await async_client.post(
        "/api/v1/resumes/cv/upload",
        files=docx_files,
        headers=DEMO_HEADERS,
    )
    assert res_docx.status_code == 201
    docx_data = res_docx.json()
    assert docx_data["file_type"] == "docx"
    assert "mayank.docx@example.com" in docx_data.get("email", "")
    assert "FastAPI" in docx_data.get("skills", [])

    # 3. Test Confirmation of extracted data
    res_confirm = await async_client.put(
        "/api/v1/resumes/cv/confirm",
        headers=DEMO_HEADERS,
        json={
            "full_name": "Mayank Pawar",
            "email": "mayank.confirmed@example.com",
            "skills": ["Python", "FastAPI", "PostgreSQL", "Docker", "Flutter"],
        },
    )
    assert res_confirm.status_code == 200
    confirmed = res_confirm.json()
    assert confirmed["is_confirmed"] is True
    assert confirmed["email"] == "mayank.confirmed@example.com"

